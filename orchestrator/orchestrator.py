"""
===============================================================
agentsHQ ORCHESTRATOR v2.0
===============================================================
Owner: Boubacar Barry — Catalyst Works Consulting
System: agentsHQ — Self-hosted multi-agent intelligence

This is the main FastAPI service. It receives tasks from:
  - Telegram (primary channel)
  - Direct HTTP POST (from Cursor, Claude Code, any tool)
  - n8n workflows (automated triggers)

Architecture:
  1. Request arrives at /run
  2. Router classifies the task type
  3. Orchestrator assembles the right crew
  4. Crew executes autonomously
  5. Output is saved, memory is updated, result is returned

See AGENTS.md for the full system architecture.
See CLAUDE.md for development guidelines.
===============================================================
"""

import os
import json
import asyncio
import threading
import uuid
import logging
import time
from datetime import datetime
from typing import Optional

import httpx
from fastapi import FastAPI, HTTPException, BackgroundTasks, Request, Depends, Header, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

# Task types that produce a tangible saved artifact (HTML, PDF, report, email template, etc.).
# Only these task types trigger a save to GitHub and Google Drive.
# Query-only tasks (CRM lookups, chat, notion reads, KPI checks) never save to Drive/GitHub.
SAVE_REQUIRED_TASK_TYPES = {
    "research_report",
    "consulting_deliverable",
    "website_build",
    "app_build",
    "3d_website_build",
    "code_task",
    "general_writing",
    "social_content",
    "linkedin_x_campaign",
    "voice_polishing",
    "crm_outreach",
    "hunter_task",
    "content_push_to_drive",
    "skill_build",
}

# Task types that benefit from pre-task memory recall.
# Simple/single-pass tasks (gws, hunter, social) are excluded — no benefit, adds latency.
MEMORY_GATED_TASK_TYPES = {
    "research_report",
    "consulting_deliverable",
    "website_build",
    "web_builder",
    "3d_web_builder",
    "notion_architect",
    "copywriting",
    "cold_outreach",
    "email_draft",
}

# In-memory tracker: chat_id -> last completed job metadata
# Used by praise/critique detector to pair feedback with prior output.
_last_completed_job: dict = {}

# In-memory project context: chat_id -> active project name
# Set via /switch <project-name>; used as session_key prefix for crews
_active_project: dict = {}

# OpenSpace Evolution Hook
try:
    from skills.openspace_skill.openspace_tool import openspace_tool
except ImportError:
    openspace_tool = None  # type: ignore[assignment]

# Catalyst Daily Ignition Scheduler
try:
    from scheduler import start_scheduler
except ImportError:
    start_scheduler = None

# Configure logging
LOG_DIR = os.environ.get("AGENTS_LOG_DIR", "/app/logs")
os.makedirs(LOG_DIR, exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler(os.path.join(LOG_DIR, "orchestrator.log"), mode="a")
    ]
)
logger = logging.getLogger("agentsHQ")

# ── App setup ──────────────────────────────────────────────────
app = FastAPI(
    title="agentsHQ Orchestrator",
    description="Self-hosted multi-agent intelligence for Catalyst Works Consulting",
    version="2.0.0"
)

# CORS narrowed to known origins. Browser chat uses the same domain, so no CORS
# is strictly required in prod -- but keeping localhost for dev iterations + the
# prod domain as explicit allowlist (was "*", which exposed XSRF surface on
# authenticated routes).
_CORS_ALLOWED = [
    "https://agentshq.boubacarbarry.com",
    "http://localhost:3000",
    "http://localhost:5173",
    "http://localhost:8000",
    "http://127.0.0.1:3000",
    "http://127.0.0.1:5173",
    "http://127.0.0.1:8000",
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=_CORS_ALLOWED,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type", "Authorization", "X-Api-Key", "X-Internal-Token"],
    allow_credentials=True,
)

def verify_api_key(x_api_key: str = Header(default=None), authorization: Optional[str] = Header(default=None)):
    """
    Reject requests without valid auth. Fail-closed when ORCHESTRATOR_API_KEY
    is unset. Accepts two forms:
      1. X-Api-Key: <raw key>              Telegram, n8n, existing integrations
      2. Authorization: Bearer <jwt>       browser chat UI

    Dev override: set DEBUG_NO_AUTH=true to bypass. Only use locally.
    """
    expected = os.environ.get("ORCHESTRATOR_API_KEY", "")
    if not expected:
        if os.environ.get("DEBUG_NO_AUTH", "false").lower() == "true":
            return
        logger.error("verify_api_key: ORCHESTRATOR_API_KEY not configured (fail-closed)")
        raise HTTPException(status_code=500, detail="Server auth misconfigured.")

    # Raw API key header (primary path)
    if x_api_key == expected:
        return

    # Bearer JWT (browser chat) or raw key as Bearer
    if authorization and authorization.startswith("Bearer "):
        import jwt as pyjwt
        token = authorization[7:]
        if token == expected:
            return
        try:
            pyjwt.decode(token, expected, algorithms=["HS256"])
            return
        except pyjwt.ExpiredSignatureError:
            raise HTTPException(status_code=401, detail="Session expired. Refresh to get a new token.")
        except pyjwt.InvalidTokenError:
            pass

    raise HTTPException(status_code=401, detail="Invalid or missing auth. Use X-Api-Key or Authorization: Bearer <token>.")

@app.on_event("startup")
async def startup_event():
    """Run at service startup."""
    # Per-call LLM ledger: register litellm callback so every completion
    # logs a row to llm_calls. Must run before any crew or council fires.
    try:
        from usage_logger import install_litellm_callback
        install_litellm_callback()
    except Exception as e:
        logger.warning(f"usage_logger startup failed (non-fatal): {e}")

    if start_scheduler:
        start_scheduler()
        logger.info("Catalyst Daily Ignition initiated.")

    # Start Telegram Polling in background
    asyncio.create_task(telegram_polling_loop())
    logger.info("Telegram Polling Loop scheduled.")

# ── Request/Response models ────────────────────────────────────
class TaskRequest(BaseModel):
    task: str
    from_number: str = "unknown"
    session_key: str = "default"
    context: Optional[dict] = None  # optional extra context from caller
    callback_url: Optional[str] = None  # webhook URL to POST result when async job completes
    file_id: Optional[str] = None   # ID from /upload — orchestrator prepends extracted text
    source: Optional[str] = None    # "browser" | "telegram" | "api"

class TaskResponse(BaseModel):
    success: bool
    result: str
    task_type: str = "unknown"
    files_created: list = []
    execution_time: float = 0.0
    title: str = ""
    deliverable: str = ""

class TeamTaskRequest(BaseModel):
    subtasks: list           # [{"crew_type": str, "task": str, "label": str}]
    original_request: str
    from_number: str = "unknown"
    session_key: str = "default"

class StatusResponse(BaseModel):
    status: str
    service: str
    version: str
    task_types: list
    agents: list
    uptime_seconds: float

class AsyncTaskResponse(BaseModel):
    job_id: str
    status: str = "pending"
    message: str = "Job queued. Poll /status/{job_id} for updates."

class JobStatusResponse(BaseModel):
    job_id: str
    status: str          # pending | running | completed | failed
    task_type: str = ""
    result: str = ""
    files_created: list = []
    execution_time: float = 0.0
    error: str = ""

# Track service start time
_start_time = time.time()

# Git lock — all local git operations must acquire this before running.
# Prevents concurrent agents from corrupting the working tree.
_git_lock = threading.Lock()


# ══════════════════════════════════════════════════════════════
# CORE ORCHESTRATION LOGIC
# ══════════════════════════════════════════════════════════════

def _query_system() -> str:
    """
    Live system introspection tool. Called by the chat LLM when the user
    asks about agents, task types, system config, or capabilities.
    Always returns accurate data from the running modules — never stale.
    """
    lines = ["=== agentsHQ LIVE SYSTEM STATE ===\n"]

    # Agent registry — read from agents.py builder functions
    try:
        import agents as agent_module
        builders = [f for f in dir(agent_module) if f.startswith("build_") and f.endswith("_agent")]
        lines.append(f"AGENTS ({len(builders)} registered):")
        agent_descriptions = {
            "build_planner_agent":       "Plans and structures every task before execution",
            "build_researcher_agent":    "Finds and synthesizes information using web search",
            "build_copywriter_agent":    "Writes reports, articles, documents, and long-form copy",
            "build_social_media_agent":  "Creates posts and content in Boubacar's voice (X, LinkedIn, etc.)",
            "build_consulting_agent":    "Produces frameworks, proposals, diagnostics, strategy briefs",
            "build_web_builder_agent":   "Builds complete single-file HTML/CSS/JS websites",
            "build_app_builder_agent":   "Builds interactive web applications",
            "build_code_agent":          "Writes, debugs, and explains code in any language",
            "build_qa_agent":            "Reviews all deliverables, fixes issues, ensures professional quality",
            "build_orchestrator_agent":  "Handles unknown/ambiguous requests, escalates or improvises",
            "build_agent_creator_agent": "Designs specs for new specialist agents when a gap is identified",
        }
        for b in sorted(builders):
            name = b.replace("build_", "").replace("_agent", "").replace("_", " ").title()
            desc = agent_descriptions.get(b, "Specialist agent")
            lines.append(f"  - {name}: {desc}")
    except Exception as e:
        lines.append(f"  [Could not load agent list: {e}]")

    # Task types — read live from router
    try:
        from router import TASK_TYPES
        lines.append(f"\nTASK TYPES ({len(TASK_TYPES) - 1} actionable):")
        for key, meta in TASK_TYPES.items():
            if key == "chat":
                continue
            lines.append(f"  - {key}: {meta['description']}")
            lines.append(f"    trigger keywords: {', '.join(meta.get('keywords', [])[:5])}")
    except Exception as e:
        lines.append(f"  [Could not load task types: {e}]")

    # Recent outputs
    try:
        output_dir = os.environ.get("AGENTS_OUTPUT_DIR", "/app/outputs")
        if os.path.exists(output_dir):
            files = sorted(
                [f for f in os.listdir(output_dir) if not os.path.isdir(os.path.join(output_dir, f))],
                key=lambda f: os.path.getmtime(os.path.join(output_dir, f)),
                reverse=True
            )[:10]
            if files:
                lines.append(f"\nRECENT OUTPUT FILES (last 10 in {output_dir}):")
                for f in files:
                    lines.append(f"  - {f}")
            else:
                lines.append(f"\nOUTPUT DIRECTORY: {output_dir} (empty)")
    except Exception:
        pass

    # Infrastructure
    lines.append("\nINFRASTRUCTURE:")
    lines.append("  VPS: 72.60.209.109 — orchestrator on port 8000")
    lines.append("  Telegram bot: @agentsHQ4Bou_bot")
    lines.append("  n8n: https://n8n.srv1040886.hstgr.cloud")
    lines.append("  GitHub: https://github.com/bokar83/agentHQ")
    lines.append("  Memory: Qdrant (vector) + PostgreSQL (conversation history)")

    return "\n".join(lines)


def run_chat(message: str, session_key: str = "default") -> dict:
    """
    Direct conversational response — no crew, no tasks.
    Uses the last 10 turns of session history so the bot remembers
    everything discussed. Fast (single LLM call, ~2-3 seconds).
    """
    start_time = datetime.now()

    # Load conversation history
    history_messages = []
    try:
        from memory import get_conversation_history
        history = get_conversation_history(session_key, limit=10)
        for turn in history:
            role = turn["role"] if turn["role"] in ("user", "assistant") else "user"
            history_messages.append({"role": role, "content": turn["content"]})
        # Anthropic requires the last message to be from the user.
        # Strip trailing assistant messages so the current user message is always last.
        while history_messages and history_messages[-1]["role"] == "assistant":
            history_messages.pop()
    except Exception as e:
        logger.warning(f"Chat history load failed (non-fatal): {e}")

    # Build system prompt
    system_prompt = """You are Boubacar's personal AI assistant, built into agentsHQ.
You know Boubacar well — he is the founder of Catalyst Works Consulting, a strategic
consulting firm. He works across AI, business development, and building systems.

PERSONALITY:
- Sarcastic, witty, and fun — think a brilliant friend who roasts you a little but
  clearly has your back. Like if Bart Simpson grew up and got an MBA.
- Drop a Simpsons quote naturally every few messages. Not every message — just when
  the moment calls for it. Make it land in context, don't force it.
- Short and punchy. No padding. Get to the point with a smirk.
- When Boubacar says something obvious, call it out. When he does something great,
  acknowledge it with minimal fanfare and move on.
- You're not a yes-man. If something is a bad idea, say so — briefly, with humor.

SIMPSONS QUOTES — use these (and others you know) when the vibe is right:
- "I am so smart! S-M-R-T." — Homer, when something goes surprisingly well
- "Trying is the first step towards failure." — Homer, when Boubacar overthinks
- "In this house we obey the laws of thermodynamics!" — when constraints come up
- "It's a perfectly cromulent word." — when something unconventional works
- "Mmm... [relevant thing]" — Homer drooling format, for anything exciting
- "Don't have a cow, man." — Bart, when Boubacar stress-tests something
- "Excellent." — Mr. Burns, when a plan comes together

MEMORY:
You have memory of past conversations. Refer to it naturally when relevant.
No need to announce "based on our history" — just use it the way a friend would.

FILE RETRIEVAL:
You have a retrieve_output_file tool. Use it immediately when Boubacar asks to see,
read, get, or retrieve a file the agents created. Do NOT say "let me grab that" and
stop — call the tool and include the full content + Drive link in your reply.

TASKS:
When Boubacar asks you to do real work (write, rewrite, research, build, draft, tweet,
post, email, leads, voice, analyze, ideas, anything that needs execution), call the
forward_to_crew tool immediately with his exact message. Do not answer it yourself.
Do not explain. Just forward it. You are a pipe for work, not the worker.

CRITICAL RULE — NO EXCEPTIONS:
If you cannot fulfill a request yourself using your available tools, do NOT:
- Explain that you lack access
- List what you can't do
- Ask the user to choose between options
- Suggest manual workarounds
- Escalate to owner

Instead: call forward_to_crew immediately with the user's exact message.
The orchestrator has GWS CLI, Gmail, Notion, CRM, and many other capabilities.
Always assume the orchestrator CAN handle it. Your job is to route, not to gatekeep.

You handle directly: greetings, memory questions, file retrieval, quick factual Q&A,
and system status. Everything else — without exception — goes to the crew."""

    # Tool definitions
    tools = [
        {
            "type": "function",
            "function": {
                "name": "query_system",
                "description": (
                    "Query the live agentsHQ system state. Call this when the user asks about: "
                    "what agents exist, what task types are available, what the system can do, "
                    "recent output files, infrastructure details, or any question about the "
                    "system's current configuration and capabilities."
                ),
                "parameters": {"type": "object", "properties": {}, "required": []}
            }
        },
        {
            "type": "function",
            "function": {
                "name": "retrieve_output_file",
                "description": (
                    "Read and return the full content of a saved output file, plus its Google Drive link. "
                    "Call this when the user asks to see, read, retrieve, or get the content of a file "
                    "that was previously created by the agents. Also call this when the user asks for "
                    "the Drive link or GitHub link for a specific output file."
                ),
                "parameters": {
                    "type": "object",
                    "properties": {
                        "filename_hint": {
                            "type": "string",
                            "description": "Partial filename or keywords to match. E.g. 'disruptive ai startups' or '5-most-disruptive'."
                        }
                    },
                    "required": ["filename_hint"]
                }
            }
        },
        {
            "type": "function",
            "function": {
                "name": "save_memory",
                "description": (
                    "Persist a fact, preference, or note to long-term memory. "
                    "Call this when the user says 'remember this', 'add to memory', "
                    "'save this', 'store this', 'note this down', or shares information "
                    "they want kept (brand colors, preferences, facts, etc.). "
                    "Do NOT call for ideas — use the crew for that."
                ),
                "parameters": {
                    "type": "object",
                    "properties": {
                        "fact": {
                            "type": "string",
                            "description": "The fact or preference to save, written as a complete sentence."
                        },
                        "category": {
                            "type": "string",
                            "description": "Category label e.g. 'brand', 'preference', 'contact', 'system'."
                        }
                    },
                    "required": ["fact"]
                }
            }
        },
        {
            "type": "function",
            "function": {
                "name": "forward_to_crew",
                "description": (
                    "Forward a task to the agentsHQ crew for execution. "
                    "Call this for ANY real work: writing, rewriting, researching, building, "
                    "voice matching, tweet polishing, post drafting, email drafting, "
                    "lead queries, ideas capture, or anything requiring agents. "
                    "Pass the user's exact original message as task_text."
                ),
                "parameters": {
                    "type": "object",
                    "properties": {
                        "task_text": {
                            "type": "string",
                            "description": "The user's exact message to forward to the crew."
                        }
                    },
                    "required": ["task_text"]
                }
            }
        },
    ]

    def _retrieve_output_file(filename_hint: str) -> str:
        """Find a matching output file and return its content + Drive link."""
        output_dir = os.environ.get("AGENTS_OUTPUT_DIR", "/app/outputs")
        all_files = []
        try:
            for root, dirs, files in os.walk(output_dir):
                for f in files:
                    all_files.append(os.path.join(root, f))
        except Exception as e:
            return f"Could not scan output directory: {e}"

        if not all_files:
            return "No output files found in the outputs directory."

        # Score files by hint keyword match
        hint_lower = filename_hint.lower()
        hint_words = hint_lower.replace("-", " ").replace("_", " ").split()
        scored = []
        for fp in all_files:
            name = os.path.basename(fp).lower().replace("-", " ").replace("_", " ")
            score = sum(1 for w in hint_words if w in name)
            scored.append((score, os.path.getmtime(fp), fp))

        scored.sort(key=lambda x: (x[0], x[1]), reverse=True)
        best_score, _, best_path = scored[0]

        if best_score == 0:
            # Fall back to most recent file
            all_files.sort(key=os.path.getmtime, reverse=True)
            best_path = all_files[0]

        try:
            with open(best_path, "r", encoding="utf-8") as f:
                content = f.read()
        except Exception as e:
            return f"Found file {os.path.basename(best_path)} but could not read it: {e}"

        # Try to get Drive link by uploading
        drive_url = None
        try:
            from saver import save_to_drive
            title = os.path.basename(best_path).replace("-", " ").replace("_", " ")
            drive_url = save_to_drive(title, "research_report", content)
        except Exception:
            pass

        result_parts = [f"File: {os.path.basename(best_path)}", ""]
        if drive_url:
            result_parts.append(f"Drive link: {drive_url}")
            result_parts.append("")
        result_parts.append(content[:3000])
        if len(content) > 3000:
            result_parts.append("\n[... content truncated — reply 'more' for the rest]")
        return "\n".join(result_parts)

    # Assemble messages: system + history + current message
    messages = [{"role": "system", "content": system_prompt}]
    messages.extend(history_messages)
    messages.append({"role": "user", "content": message})

    try:
        import openai
        client = openai.OpenAI(
            api_key=os.environ.get("OPENROUTER_API_KEY"),
            base_url="https://openrouter.ai/api/v1",
            default_headers={
                "HTTP-Referer": "https://agentshq.catalystworks.com",
                "X-Title": "agentsHQ Chat"
            }
        )

        response = client.chat.completions.create(
            model="anthropic/claude-haiku-4.5",
            messages=messages,
            tools=tools,
            tool_choice="auto",
            temperature=0.85,
        )

        msg = response.choices[0].message

        # Handle tool calls
        if msg.tool_calls:
            messages.append(msg)
            for tool_call in msg.tool_calls:
                fn_name = tool_call.function.name
                if fn_name == "query_system":
                    tool_result = _query_system()
                    logger.info("Chat used query_system tool")
                elif fn_name == "retrieve_output_file":
                    args = json.loads(tool_call.function.arguments or "{}")
                    tool_result = _retrieve_output_file(args.get("filename_hint", ""))
                    logger.info(f"Chat used retrieve_output_file: {args.get('filename_hint')}")
                elif fn_name == "save_memory":
                    args = json.loads(tool_call.function.arguments or "{}")
                    fact = args.get("fact", "")
                    category = args.get("category", "general")
                    try:
                        from memory import save_conversation_turn, save_to_memory
                        tag = f"[MEMORY:{category.upper()}] {fact}"
                        save_conversation_turn(session_key, "assistant", tag)
                        save_to_memory(
                            task_request=tag,
                            task_type="memory_capture",
                            result_summary=fact,
                            files_created=[],
                            execution_time=0,
                            from_number=session_key,
                        )
                        logger.info(f"Chat used save_memory: {fact[:80]}")
                        tool_result = f"Saved to memory: {fact}"
                    except Exception as mem_e:
                        tool_result = f"Memory save failed: {mem_e}"
                elif fn_name == "forward_to_crew":
                    args = json.loads(tool_call.function.arguments or "{}")
                    task_text = args.get("task_text", message)
                    try:
                        fwd_result = run_orchestrator(
                            task_request=task_text,
                            from_number=session_key,
                            session_key=session_key
                        )
                        tool_result = fwd_result.get("result") or fwd_result.get("deliverable") or "Crew completed the task."
                        logger.info(f"Chat forwarded to crew: {task_text[:60]}")
                    except Exception as fwd_e:
                        tool_result = f"Crew error: {fwd_e}"
                        logger.error(f"forward_to_crew failed: {fwd_e}")
                else:
                    tool_result = "Unknown tool."
                messages.append({
                    "role": "tool",
                    "tool_call_id": tool_call.id,
                    "content": tool_result
                })
            followup = client.chat.completions.create(
                model="anthropic/claude-haiku-4.5",
                messages=messages,
                temperature=0.85,
            )
            reply = (followup.choices[0].message.content or "").strip()
        else:
            reply = (msg.content or "").strip()

    except Exception as e:
        logger.error(f"Chat LLM call failed: {e}")
        reply = "Sorry, I hit an error. Try again in a moment."

    # Save this exchange to history
    try:
        from memory import save_conversation_turn
        save_conversation_turn(session_key, "user", message)
        save_conversation_turn(session_key, "assistant", reply)
    except Exception as e:
        logger.warning(f"Chat history save failed (non-fatal): {e}")

    execution_time = (datetime.now() - start_time).total_seconds()
    logger.info(f"Chat response for session '{session_key}' in {execution_time:.1f}s")

    return {
        "success": True,
        "result": reply,
        "full_output": reply,
        "task_type": "chat",
        "files_created": [],
        "execution_time": execution_time,
        "classification": {"task_type": "chat", "confidence": 1.0, "is_unknown": False}
    }


def run_orchestrator(task_request: str, from_number: str = "unknown", session_key: str = "default") -> dict:
    """
    Main orchestration function.

    1. Load session history for context
    2. Route: classify the task type
    3. Assemble: build the right crew
    4. Execute: run the crew autonomously
    5. Save: store output to memory + conversation history
    6. Return: structured result for the caller
    """
    start_time = datetime.now()

    # Step 1: Load recent conversation history for this session
    today = start_time.strftime("%B %d, %Y")
    enriched_task = f"[Today's date: {today}. All research, recommendations, and references should reflect the current state as of {today}.]\n\n{task_request}"
    try:
        from memory import get_conversation_history
        history = get_conversation_history(session_key, limit=6)
        if history:
            history_text = "\n".join(
                f"[{h['role'].upper()}]: {h['content'][:800]}"
                for h in history
            )
            enriched_task = (
                f"--- CONVERSATION HISTORY (most recent first) ---\n"
                f"{history_text}\n"
                f"--- END HISTORY ---\n\n"
                f"CURRENT REQUEST: {task_request}"
            )
            logger.info(f"Injected {len(history)} history entries for session '{session_key}'")
    except Exception as e:
        logger.warning(f"History injection failed (non-fatal): {e}")

    # Step 2: Route
    from router import classify_task, get_crew_type, TASK_TYPES
    classification = classify_task(task_request)  # classify on raw task, not enriched
    task_type = classification.get("task_type", "unknown")
    is_unknown = classification.get("is_unknown", False)

    logger.info(f"Task classified as '{task_type}' (confidence: {classification.get('confidence', 0)})")

    # Step 0b: Pre-task memory recall (injected now that we know task_type)
    if (
        os.environ.get("MEMORY_LEARNING_ENABLED", "false").lower() == "true"
        and task_type in MEMORY_GATED_TASK_TYPES
    ):
        try:
            from memory import query_memory
            memory_lines = []

            past_work = query_memory(task_request, top_k=3)
            if past_work:
                memory_lines.append("--- RELEVANT PAST WORK ---")
                for pw in past_work:
                    memory_lines.append(
                        f"- [{pw.get('task_type','?')}] {pw.get('summary','')[:200]} (date: {pw.get('date','?')})"
                    )

            past_lessons = query_memory(task_request, top_k=5, collection="agentshq_learnings")
            positive = [l for l in past_lessons if l.get("lesson_type") == "positive"]
            negative = [l for l in past_lessons if l.get("lesson_type") == "negative"]
            if positive:
                memory_lines.append("--- WHAT WORKED WELL FOR THIS TASK TYPE ---")
                for l in positive:
                    memory_lines.append(f"- {l.get('extracted_pattern','')[:200]}")
            if negative:
                memory_lines.append("--- WHAT TO AVOID FOR THIS TASK TYPE ---")
                for l in negative:
                    memory_lines.append(f"- {l.get('extracted_pattern','')[:200]}")

            if memory_lines:
                memory_block = "\n".join(memory_lines) + "\n--- END MEMORY ---\n\n"
                # Cap total enriched_task at 6000 chars to prevent context window overflow
                combined = memory_block + enriched_task
                if len(combined) > 6000:
                    allowed = max(0, 6000 - len(enriched_task))
                    memory_block = memory_block[:allowed]
                enriched_task = memory_block + enriched_task
                logger.info(f"Memory recall: {len(past_work)} past tasks + {len(past_lessons)} lessons injected for {task_type}")
        except Exception as e:
            logger.warning(f"Memory recall failed (non-fatal): {e}")

    # Step 3: Direct dispatch for crm_outreach (no crew -- pure Supabase + Gmail API)
    if task_type == "crm_outreach":
        try:
            from skills.outreach.outreach_tool import run_outreach
            contact_all = "contact all" in task_request.lower() or "all leads" in task_request.lower()
            outreach_result = run_outreach(contact_all=contact_all)
            drafted = outreach_result.get("drafted", 0)
            skipped = outreach_result.get("skipped", 0)
            results = outreach_result.get("results", [])
            error = outreach_result.get("error")

            if error:
                deliverable = f"Outreach failed: {error}"
            elif drafted == 0:
                deliverable = (
                    "No drafts created. Either no leads have confirmed emails yet, "
                    "or all eligible leads have already been contacted. "
                    "Run 'find leads' to add new ones."
                )
            else:
                lines = [f"Cold outreach drafts created in boubacar@catalystworks.consulting ({drafted} drafts):\n"]
                for r in results:
                    if r.get("status") == "drafted":
                        lines.append(f"- {r['name']} | {r['company']} | {r['email']} | {r['subject']}")
                if skipped:
                    lines.append(f"\n{skipped} lead(s) failed (Gmail API error -- check logs).")
                deliverable = "\n".join(lines)

            return {
                "success": True,
                "task_type": task_type,
                "deliverable": deliverable,
                "execution_time": (datetime.now() - start_time).total_seconds(),
            }
        except Exception as e:
            logger.error(f"crm_outreach direct dispatch failed: {e}")
            return {"success": False, "task_type": task_type, "deliverable": f"Outreach error: {e}", "execution_time": 0}

    # Step 3a: Direct dispatch for research_report: bypass CrewAI entirely.
    # CrewAI's max_iter fallback triggers Anthropic's "assistant message prefill"
    # 400 on heavy research prompts. See docs/superpowers/plans/2026-04-20-research-engine-bypass.md.
    if task_type == "research_report":
        try:
            from research_engine import run_research
            research_result = run_research(user_prompt=enriched_task)
            deliverable = research_result.get("answer") or ""
            execution_time = (datetime.now() - start_time).total_seconds()

            if not research_result.get("success"):
                err = research_result.get("error", "unknown")
                logger.error(f"research_engine failed: {err}")
                deliverable = (
                    "Research couldn't complete. Try narrowing the request to one question "
                    f"or one zip code, then ask again. (Diagnostic: {err})"
                )

            try:
                from memory import save_to_memory, save_conversation_turn
                save_to_memory(
                    task_request=task_request,
                    task_type=task_type,
                    result_summary=deliverable[:1000],
                    files_created=[],
                    execution_time=execution_time,
                    from_number=from_number,
                )
                save_conversation_turn(session_key, "user", task_request)
                save_conversation_turn(session_key, "assistant", deliverable[:1000])
            except Exception as mem_err:
                logger.warning(f"Memory save failed (non-fatal): {mem_err}")

            return {
                "success": True,
                "result": deliverable,
                "task_type": task_type,
                "files_created": [],
                "execution_time": execution_time,
                "title": task_request[:80].strip(),
                "deliverable": deliverable,
            }
        except Exception as e:
            logger.error(f"research_engine dispatch failed, falling back to CrewAI: {e}", exc_info=True)
            # fall through to CrewAI crew below

    # Step 3: Assemble crew
    from crews import assemble_crew

    if is_unknown:
        crew_type = "unknown_crew"
    else:
        crew_type = get_crew_type(task_type) or "unknown_crew"

    crew = assemble_crew(crew_type, enriched_task)

    # Step 4: Execute
    # Tag every LLM call inside this kickoff with crew/task/session metadata
    # so the per-call ledger (llm_calls) can attribute spend correctly.
    try:
        from usage_logger import current_call_context
        current_call_context.set({
            "project": "agentsHQ",
            "task_type": task_type,
            "crew_name": crew_type,
            "session_key": session_key,
        })
    except Exception:
        pass

    logger.info(f"Kicking off crew: {crew_type}")
    result = crew.kickoff()
    result_str = result.raw if hasattr(result, 'raw') else str(result)
    try:
        from skills.boub_voice_mastery.voice_polisher import polish_voice
        result_str = polish_voice(result_str)
    except Exception:
        pass

    # Extract deliverable — everything after "DELIVERABLE:" marker
    lower = result_str.lower()
    idx = lower.find("deliverable:")
    deliverable = result_str[idx + len("deliverable:"):].strip() if idx != -1 else result_str.strip()
    title = task_request[:80].strip()

    end_time = datetime.now()
    execution_time = (end_time - start_time).total_seconds()
    
    # Step 4: Collect files created
    files_created = []
    output_dir = os.environ.get("AGENTS_OUTPUT_DIR", "/app/outputs")
    if os.path.exists(output_dir):
        all_files = os.listdir(output_dir)
        # Get files created in this execution window
        recent_files = [
            f for f in all_files
            if os.path.getmtime(os.path.join(output_dir, f)) >= start_time.timestamp()
        ]
        files_created = recent_files
    
    # Step 5: Save to memory + conversation history
    try:
        from memory import save_to_memory, save_conversation_turn
        result_summary = result_str[:1000] if len(result_str) > 1000 else result_str
        save_to_memory(
            task_request=task_request,
            task_type=task_type,
            result_summary=result_summary,
            files_created=files_created,
            execution_time=execution_time,
            from_number=from_number
        )
        # Save this exchange to conversation history for follow-up context
        save_conversation_turn(session_key, "user", task_request)
        save_conversation_turn(session_key, "assistant", result_summary)
    except Exception as e:
        logger.warning(f"Memory save failed (non-fatal): {e}")

    # Step 6: Sync artifact to Notion (non-blocking daemon thread)
    try:
        from memory import sync_artifact_to_notion
        threading.Thread(
            target=sync_artifact_to_notion,
            args=(task_request, task_type, result_summary, files_created, execution_time, session_key),
            daemon=True
        ).start()
    except Exception:
        pass  # non-fatal

    # Build Telegram-friendly summary
    summary = _build_summary(task_type, result_str, files_created, execution_time)

    # Save overflow so user can reply 'more' if output was truncated
    _save_overflow_if_needed(session_key, result_str, task_type)

    return {
        "success": True,
        "result": summary,
        "task_type": task_type,
        "files_created": files_created,
        "execution_time": execution_time,
        "title": title,
        "deliverable": deliverable,
    }


def run_team_orchestrator(subtasks: list, original_request: str, from_number: str = "unknown") -> dict:
    """
    Run multiple crews in parallel and synthesize results.
    Called by POST /run-team.

    subtasks: [{"crew_type": str, "task": str, "label": str}, ...]
    """
    from crews import run_parallel_team, build_team_synthesis_crew

    start_time = datetime.now()

    # Phase 1: run all subtasks in parallel
    teammate_results = run_parallel_team(subtasks)

    successful = [r for r in teammate_results if r["success"]]
    failed     = [r for r in teammate_results if not r["success"]]

    if failed:
        logger.warning(f"[agent-team] {len(failed)} teammate(s) failed: {[f['label'] for f in failed]}")

    if not successful:
        raise RuntimeError("All teammates failed — no results to synthesize.")

    # Phase 2: synthesize
    synthesis_crew = build_team_synthesis_crew(original_request, successful)
    final_result   = synthesis_crew.kickoff()
    result_str     = final_result.raw if hasattr(final_result, 'raw') else str(final_result)

    execution_time = (datetime.now() - start_time).total_seconds()

    # Collect files created
    files_created = []
    output_dir = os.environ.get("AGENTS_OUTPUT_DIR", "/app/outputs")
    if os.path.exists(output_dir):
        files_created = [
            f for f in os.listdir(output_dir)
            if os.path.getmtime(os.path.join(output_dir, f)) >= start_time.timestamp()
        ]

    # Save to memory
    try:
        from memory import save_to_memory
        save_to_memory(
            task_request=original_request,
            task_type="agent_team",
            result_summary=result_str[:1000],
            files_created=files_created,
            execution_time=execution_time,
            from_number=from_number
        )
    except Exception as e:
        logger.warning(f"Memory save failed (non-fatal): {e}")

    summary = _build_summary("agent_team", result_str, files_created, execution_time)

    return {
        "success": True,
        "result": summary,
        "full_output": result_str,
        "task_type": "agent_team",
        "teammate_count": len(subtasks),
        "teammates_succeeded": len(successful),
        "teammates_failed": len(failed),
        "files_created": files_created,
        "execution_time": execution_time,
    }


def _build_summary(
    task_type: str,
    full_output: str,
    files_created: list,
    execution_time: float
) -> str:
    """
    Build a Telegram message with the full deliverable content.
    Always includes: what was done, time taken, files saved, and full output.
    """

    type_labels = {
        "agent_team":              "Team task complete",
        "website_build":           "Website built",
        "app_build":               "App built",
        "research_report":         "Research complete",
        "consulting_deliverable":  "Consulting deliverable ready",
        "social_content":          "Social content created",
        "linkedin_x_campaign":     "LinkedIn/X campaign ready",
        "code_task":               "Code task complete",
        "general_writing":         "Document ready",
        "agent_creation":          "Agent proposal submitted",
        "gws_task":                "Google Workspace task complete",
        "unknown":                 "Task complete",
    }

    label = type_labels.get(task_type, "Task complete")

    lines = [
        f"--- {label} ---",
        f"Time: {execution_time:.0f}s",
    ]

    if files_created:
        lines.append(f"Files saved: {', '.join(files_created)}")

    lines.append("")  # blank line before content

    # Send full output up to Telegram's 4096 char limit
    # Reserve ~150 chars for the header lines above
    MAX_CONTENT = 3700
    if full_output and len(full_output) > 50:
        content = full_output.strip()
        if len(content) > MAX_CONTENT:
            content = content[:MAX_CONTENT] + "\n\n[reply 'more' to see the rest]"
        lines.append(content)
    else:
        lines.append("[No output content returned by agents]")

    return "\n".join(lines)


def _save_overflow_if_needed(session_key: str, full_output: str, task_type: str) -> None:
    """Save overflow to DB if output exceeds one Telegram message."""
    MAX_CONTENT = 3700
    if full_output and len(full_output.strip()) > MAX_CONTENT:
        try:
            from memory import save_overflow
            save_overflow(session_key, full_output.strip(), MAX_CONTENT, task_type)
        except Exception as e:
            logger.warning(f"Overflow save failed (non-fatal): {e}")


def _run_background_job(
    task: str,
    from_number: str,
    session_key: str,
    job_id: str,
    classification: dict = None,
) -> None:
    """
    Background worker: runs the crew, sends progress pings, saves output,
    and delivers the final result to Telegram.
    Called via FastAPI BackgroundTasks — runs in a thread pool.
    """
    from notifier import send_progress_ping, send_result, send_message
    from saver import save_to_github, save_to_drive

    chat_id = from_number  # from_number IS the Telegram chat ID

    # Register job in queue so Telegram tasks are traceable
    try:
        from memory import create_job, update_job
        create_job(job_id, session_key, from_number, task)
    except Exception as e:
        logger.warning(f"Job {job_id}: could not register in queue: {e}")

    # ── Progress ping timer ──────────────────────────────────
    # First ping at 60s so user knows it's working, then every 5 min.
    # Watchdog at 10 min: if still running, send an error and bail.
    _stop_ping = threading.Event()

    def _ping_loop():
        # First ping after 60s
        if _stop_ping.wait(60):
            return
        send_progress_ping(chat_id)
        # Subsequent pings every 5 min, watchdog at 10 min total (1 more ping)
        if _stop_ping.wait(300):
            return
        send_progress_ping(chat_id)
        # Watchdog: if still running after ~10 min total, something is hung
        if not _stop_ping.wait(300):
            send_message(chat_id, "⚠️ Task is taking longer than expected (10+ min). It may be stuck — please try again or check with the team.")
            _stop_ping.set()

    ping_thread = threading.Thread(target=_ping_loop, daemon=True)
    ping_thread.start()

    result = {}  # ensure result is always defined before try/finally
    try:
        # ── Run the crew ─────────────────────────────────────
        result = run_orchestrator(
            task_request=task,
            from_number=from_number,
            session_key=session_key,
        )

        summary     = result.get("result", "")
        task_type   = result.get("task_type", "unknown")
        title       = result.get("title", task[:80])
        deliverable = result.get("deliverable", summary)

        # ── Save (only for tasks that produce tangible artifacts) ────────────
        github_url = ""
        drive_url  = ""
        if task_type in SAVE_REQUIRED_TASK_TYPES:
            github_url = save_to_github(title, task_type, deliverable)
            drive_url  = save_to_drive(title, task_type, deliverable)
        else:
            logger.info(f"Job {job_id}: task_type '{task_type}' is query-only — skipping Drive/GitHub save")

        # ── Deliver ──────────────────────────────────────────
        send_result(chat_id, summary, drive_url, github_url)
        try:
            update_job(job_id, status="completed", result=summary)
        except Exception as e:
            logger.warning(f"Job {job_id}: could not update job status to completed: {e}")

        # ── Record last completed job for feedback detection ──────────────────
        _last_completed_job[chat_id] = {
            "job_id": job_id,
            "task_type": task_type,
            "task_request": task,
            "result_summary": summary[:1000],
            "delivered_at": datetime.utcnow(),
        }

        # ── Compound request: email follow-up ────────────────────
        # If the original message asked to "also send me an email about this",
        # spin a minimal GWS crew to draft the summary email after the main task.
        _has_email_followup = (classification or {}).get("has_email_followup", False)
        if _has_email_followup and task_type not in ("chat", "gws_task", "crm_outreach"):
            try:
                from crews import build_gws_crew
                email_task_text = (
                    f"Create a Gmail draft to bokar83@gmail.com, catalystworks.ai@gmail.com, and boubacar@catalystworks.consulting "
                    f"with subject 'agentsHQ Result: {title[:60]}'. "
                    f"Body: Format the following result as clean HTML with a header, "
                    f"bullet points for key findings, and a closing note. "
                    f"Result to format:\n\n{summary[:3000]}"
                )
                email_crew = build_gws_crew(email_task_text)
                email_crew.kickoff()
                send_message(chat_id, "Email draft created in your Gmail — go check it.")
                logger.info(f"Compound email follow-up draft created for job {job_id}")
            except Exception as e:
                logger.warning(f"Compound email follow-up failed (non-fatal): {e}")

        # ── Email hunter results ──────────────────────────────
        if task_type == "hunter_task":
            try:
                from notifier import send_hunter_report
                send_hunter_report(leads_output=deliverable, scoreboard=summary)
                logger.info("Hunter report emailed to Boubacar.")
            except Exception as e:
                logger.warning(f"Hunter email failed (non-fatal): {e}")

    except Exception as e:
        logger.error(f"Background job {job_id} failed: {e}", exc_info=True)
        try:
            from notifier import send_message
            send_message(chat_id, f"Sorry — something went wrong with your task. Error: {str(e)[:200]}")
        except Exception:
            logger.critical(f"Background job {job_id}: result delivery AND error notification both failed. User received nothing.")
        try:
            update_job(job_id, status="failed", result=str(e))
        except Exception:
            pass
    finally:
        _stop_ping.set()  # cancel ping loop regardless of success/failure
        
        # ── Trigger Evolution (Self-Learning) ────────────────
        # We run this in a separate thread to avoid blocking the main background task completion
        # and to keep it completely "background" as requested.
        try:
            if result.get("success"):
                threading.Thread(
                    target=_trigger_evolution,
                    args=(task, result.get("full_output", "")),
                    daemon=True
                ).start()
        except Exception as e:
            logger.warning(f"Evolution trigger failed: {e}")

        # ── Trigger Learning Extraction ──────────────────────────────────
        # Fires only on success + MEMORY_LEARNING_ENABLED=true + gated task types
        # Runs in daemon thread — does not block result delivery
        try:
            if (
                result.get("success")
                and os.environ.get("MEMORY_LEARNING_ENABLED", "false").lower() == "true"
                and result.get("task_type", "") in MEMORY_GATED_TASK_TYPES
            ):
                from memory import extract_and_save_learnings
                threading.Thread(
                    target=extract_and_save_learnings,
                    args=(task, result.get("task_type", "unknown"), result.get("result", "")[:1000]),
                    daemon=True
                ).start()
        except Exception as e:
            logger.warning(f"Learning extraction trigger failed (non-fatal): {e}")

def _trigger_evolution(task_instruction: str, result_output: str) -> None:
    """
    Helper to run OpenSpace evolution in the background.
    Uses the tool wrapper to either evolve a skill or suggest a new one.
    """
    import asyncio
    try:
        # Check if evolution is explicitly disabled via env
        if os.environ.get("DISABLE_EVOLUTION") == "true":
            return

        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        
        # Only evolve if the task was significant (token efficiency)
        if openspace_tool is not None and (len(task_instruction) > 20 or len(result_output) > 100):
            logger.info("Triggering OpenSpace evolution...")
            loop.run_until_complete(openspace_tool.execute_async(
                command="evolve",
                task_instruction=f"Task: {task_instruction}\n\nResult:\n{result_output}"
            ))
            logger.info("OpenSpace evolution check complete.")
            
    except Exception as e:
        logger.error(f"Background evolution failed: {e}")
    finally:
        loop.close()


_TASK_KEYWORDS = [
    "write", "create", "build", "research", "analyze", "make",
    "draft", "generate", "code", "script", "website", "report",
    "proposal", "post", "email", "article",
    "find", "hunt", "leads", "prospect", "run the", "hunter",
    "news", "brief", "headlines",
    "task", "tasks", "due", "overdue", "notion", "calendar",
    "open tasks", "past due", "pending",
]
_CHAT_PREFIXES = (
    "what is my", "what's my", "how much", "do you", "can you tell",
    "hey", "hi ", "hello", "thanks", "thank you", "what did",
    "do you remember", "remind me", "what have we", "what was",
)

def _shortcut_classify(msg: str):
    """
    Run keyword shortcuts BEFORE the obvious-chat pre-filter.
    Returns a task_type string if matched, else None.
    This prevents short messages from being swallowed by _classify_obvious_chat().
    """
    from router import _keyword_shortcut
    return _keyword_shortcut(msg)


def _classify_obvious_chat(msg: str) -> bool:
    """
    Returns True only for unmistakable single-word greetings with no task content.
    Everything else goes through classify_task (LLM fallback included).
    This prevents the heuristic from swallowing short natural-language task requests.
    """
    m = msg.strip().lower().rstrip("!.,?")
    return m in {"hi", "hey", "hello", "thanks", "thank you", "morning", "good morning", "good evening"}


# ══════════════════════════════════════════════════════════════
# PRAISE / CRITIQUE DETECTION HELPERS
# ══════════════════════════════════════════════════════════════

_PRAISE_SIGNALS = {
    "good job", "great", "well done", "perfect", "excellent",
    "love it", "nice work", "brilliant", "solid", "nailed it",
    "that's great", "awesome", "fantastic", "good work", "nice",
}

_CRITIQUE_SIGNALS = {
    "wrong", "bad", "not good", "don't like", "fix", "redo",
    "wasn't", "wasn't done", "missed", "forgot", "too long",
    "too short", "off", "incorrect", "weird", "not what",
    "could be better", "needs work", "poorly", "weak",
}


def _is_praise(text: str) -> bool:
    """Return True if the message is short explicit praise."""
    t = text.lower().strip()
    if not t or len(t) > 80:
        return False
    return any(p in t for p in _PRAISE_SIGNALS)


def _is_feedback_on_prior_job(text: str, chat_id: str) -> bool:
    """
    Return True if the message looks like critique on the last completed job.
    Requires: job delivered within 60 min AND critique signal words present.
    """
    if chat_id not in _last_completed_job:
        return False
    job = _last_completed_job[chat_id]
    elapsed = (datetime.utcnow() - job["delivered_at"]).total_seconds()
    if elapsed > 3600:
        return False
    t = text.lower()
    return any(s in t for s in _CRITIQUE_SIGNALS)


# ══════════════════════════════════════════════════════════════
# API ENDPOINTS
# ══════════════════════════════════════════════════════════════

async def process_telegram_update(update: dict):
    """
    Unified processor for Telegram updates (webhook or polling).
    """
    message = update.get("message") or update.get("edited_message")
    if not message:
        return

    text = message.get("text", "").strip()
    chat_id = str(message.get("chat", {}).get("id", ""))
    sender_id = str(message.get("from", {}).get("id", ""))

    if not text or not chat_id:
        return

    # ── Sender authentication ────────────────────────────────────
    _allowed_raw = os.environ.get("ALLOWED_USER_IDS", "")
    _allowed_ids = {uid.strip() for uid in _allowed_raw.split(",") if uid.strip()}
    if _allowed_ids and sender_id not in _allowed_ids:
        from notifier import send_message as _send
        _send(chat_id, "Sorry, you are not authorised to use this bot.")
        logger.warning(f"Unauthorised Telegram access attempt from sender_id={sender_id}")
        return

    # -- Doc routing emoji + text command handlers --
    _EMOJI_COMMANDS = ("✅", "✏️", "🆕", "❌", "➕")
    _TEXT_ALIASES = {"yes": "✅", "confirm": "✅", "approved": "✅", "approve": "✅", "flag": "❌", "discard": "❌", "reject": "❌"}
    _matched_emoji = next((e for e in _EMOJI_COMMANDS if text.startswith(e)), None)
    if not _matched_emoji:
        _text_lower = text.strip().lower().split()[0] if text.strip() else ""
        _matched_emoji = _TEXT_ALIASES.get(_text_lower)
        # Also handle "edit field:value" text form -- map to ✏️
        if not _matched_emoji and _text_lower == "edit":
            _matched_emoji = "✏️"
    if _matched_emoji:
        from notifier import send_message as _send_emoji

        def _get_pending_doc(conn, reply_msg_id):
            """Return the pending doc record, using reply_to_message_id or latest unresolved."""
            cur = conn.cursor()
            if reply_msg_id:
                cur.execute(
                    "SELECT * FROM notebooklm_pending_docs WHERE telegram_message_id = %s AND resolved = false LIMIT 1",
                    (str(reply_msg_id),),
                )
            else:
                cur.execute(
                    "SELECT * FROM notebooklm_pending_docs WHERE resolved = false ORDER BY created_at DESC LIMIT 1"
                )
            row = cur.fetchone()
            if row is None:
                return None
            cols = [desc[0] for desc in cur.description]
            return dict(zip(cols, row))

        def _db_connect():
            import psycopg2
            return psycopg2.connect(
                host=os.environ.get("POSTGRES_HOST", "orc-postgres"),
                database=os.environ.get("POSTGRES_DB", "postgres"),
                user=os.environ.get("POSTGRES_USER", "postgres"),
                password=os.environ.get("POSTGRES_PASSWORD", ""),
                port=int(os.environ.get("POSTGRES_PORT", 5432)),
            )

        reply_msg_id = message.get("reply_to_message", {}).get("message_id") if message.get("reply_to_message") else None

        # ---- ✅ Confirm and file document ----
        if _matched_emoji == "✅":
            try:
                import json as _json
                from datetime import datetime as _dt
                import sys as _sys; _sys.path.insert(0, "/app/orchestrator_skills") if "/app/orchestrator_skills" not in _sys.path else None; from doc_routing.gws_cli_tools import GWSDriveMoveRenameTool, GWSSheetsAppendRowTool
                conn = _db_connect()
                record = _get_pending_doc(conn, reply_msg_id)
                if not record:
                    _send_emoji(chat_id, "No pending document found. It may have already been filed.")
                    conn.close()
                    return
                review_queue_folder_id = os.environ.get("NOTEBOOKLM_REVIEW_QUEUE_FOLDER_ID", "")
                target_path = record.get("target_folder_path", "") or ""
                if target_path.startswith("00_Review_Queue/"):
                    logger.warning(f"Doc routing: target_folder_path '{target_path}' is unresolved -- skipping Drive move for record_id={record['record_id']}")
                    new_parent_id = review_queue_folder_id
                    skipped_move = True
                else:
                    new_parent_id = review_queue_folder_id  # fallback -- path-to-ID resolution not yet implemented
                    logger.warning(f"Doc routing: Drive folder ID resolution not implemented for path '{target_path}' -- using review queue folder as fallback for record_id={record['record_id']}")
                    skipped_move = False
                if not skipped_move:
                    GWSDriveMoveRenameTool()._run(_json.dumps({
                        "file_id": record["drive_file_id"],
                        "new_name": record["standardized_filename"],
                        "new_parent_id": new_parent_id,
                        "old_parent_id": os.environ.get("NOTEBOOKLM_LIBRARY_ROOT_ID", ""),
                    }))
                queue_sheet_id = os.environ.get("NOTEBOOKLM_QUEUE_SHEET_ID", "")
                GWSSheetsAppendRowTool()._run(_json.dumps({
                    "spreadsheet_id": queue_sheet_id,
                    "range": "Sheet1!A:I",
                    "values": [[
                        record["standardized_filename"],
                        "",
                        record["domain"],
                        record["project_id"],
                        record["topic_or_client"],
                        record["doc_type"],
                        record["notebook_assignment"],
                        _dt.utcnow().strftime("%Y-%m-%d"),
                        "No",
                    ]],
                }))
                cur = conn.cursor()
                cur.execute(
                    "UPDATE notebooklm_pending_docs SET resolved = true WHERE record_id = %s",
                    (record["record_id"],),
                )
                conn.commit()
                conn.close()
                _send_emoji(
                    chat_id,
                    f"Filed: {record['standardized_filename']}\nFolder: {target_path}\nNotebook queue: updated.",
                )
            except Exception as _exc:
                logger.error(f"Emoji handler ✅ error: {_exc}", exc_info=True)
                _send_emoji(chat_id, "Error processing command. Check logs.")
            return

        # ---- ✏️ Edit a field on the pending doc ----
        if _matched_emoji == "✏️":
            try:
                payload = text[len("✏️"):].strip()
                field_map = {
                    "folder": "target_folder_path",
                    "project": "project_id",
                    "doctype": "doc_type",
                    "notebook": "notebook_assignment",
                    "name": "standardized_filename",
                }
                if ":" not in payload:
                    _send_emoji(chat_id, "Edit format: ✏️ field:value  (valid fields: folder, project, doctype, notebook, name)")
                    return
                field_key, field_val = payload.split(":", 1)
                field_key = field_key.strip().lower()
                field_val = field_val.strip()
                if field_key not in field_map:
                    _send_emoji(chat_id, f"Unknown field '{field_key}'. Valid fields: {', '.join(field_map.keys())}")
                    return
                conn = _db_connect()
                record = _get_pending_doc(conn, reply_msg_id)
                if not record:
                    _send_emoji(chat_id, "No pending document found to edit.")
                    conn.close()
                    return
                db_col = field_map[field_key]
                cur = conn.cursor()
                cur.execute(
                    f"UPDATE notebooklm_pending_docs SET {db_col} = %s WHERE record_id = %s",
                    (field_val, record["record_id"]),
                )
                conn.commit()
                # Re-fetch updated record
                cur.execute(
                    "SELECT * FROM notebooklm_pending_docs WHERE record_id = %s",
                    (record["record_id"],),
                )
                row = cur.fetchone()
                cols = [desc[0] for desc in cur.description]
                updated = dict(zip(cols, row))
                conn.close()
                _send_emoji(
                    chat_id,
                    f"Updated {field_key} to: {field_val}\n\nNew routing:\nFile: {updated['standardized_filename']}\nFolder: {updated['target_folder_path']}\nNotebook: {updated['notebook_assignment']}\n\nSend ✅ to confirm.",
                )
            except Exception as _exc:
                logger.error(f"Emoji handler ✏️ error: {_exc}", exc_info=True)
                _send_emoji(chat_id, "Error processing command. Check logs.")
            return

        # ---- 🆕 Create new project and file document ----
        if _matched_emoji == "🆕":
            try:
                import json as _json
                from datetime import datetime as _dt
                import sys as _sys; _sys.path.insert(0, "/app/orchestrator_skills") if "/app/orchestrator_skills" not in _sys.path else None; from doc_routing.gws_cli_tools import GWSDriveMoveRenameTool, GWSSheetsAppendRowTool
                payload = text[len("🆕"):].strip()
                if not payload:
                    _send_emoji(chat_id, "Usage: 🆕 <project name>")
                    return
                name_lower = payload.lower()
                if "client" in name_lower or (payload[0].isupper() and "research" not in name_lower):
                    prefix = "CL"
                elif "research" in name_lower or "topic" in name_lower:
                    prefix = "RS"
                elif "ops" in name_lower or "admin" in name_lower or "internal" in name_lower:
                    prefix = "OP"
                else:
                    prefix = "CW"
                conn = _db_connect()
                cur = conn.cursor()
                cur.execute(f"SELECT pg_advisory_lock(hashtext('{prefix}'))")
                cur.execute(
                    "SELECT COALESCE(MAX(CAST(SPLIT_PART(project_id, '-', 2) AS INTEGER)), 0) + 1 AS next_num "
                    "FROM notebooklm_pending_docs WHERE project_id LIKE %s",
                    (f"{prefix}-%",),
                )
                row = cur.fetchone()
                next_num = row[0] if row else 1
                new_project_id = f"{prefix}-{next_num:03d}"
                cur.execute(f"SELECT pg_advisory_unlock(hashtext('{prefix}'))")
                record = _get_pending_doc(conn, reply_msg_id)
                if not record:
                    conn.close()
                    _send_emoji(chat_id, "No pending document found to assign to the new project.")
                    return
                cur.execute(
                    "UPDATE notebooklm_pending_docs SET project_id = %s WHERE record_id = %s",
                    (new_project_id, record["record_id"]),
                )
                conn.commit()
                # Append new project to Registry sheet
                registry_sheet_id = os.environ.get("NOTEBOOKLM_REGISTRY_SHEET_ID", "")
                GWSSheetsAppendRowTool()._run(_json.dumps({
                    "spreadsheet_id": registry_sheet_id,
                    "range": "Projects!A:H",
                    "values": [[
                        new_project_id,
                        payload,
                        payload,
                        "",
                        "",
                        "",
                        "active",
                        _dt.utcnow().strftime("%Y-%m-%d"),
                    ]],
                }))
                # Re-fetch updated record for filing
                cur.execute(
                    "SELECT * FROM notebooklm_pending_docs WHERE record_id = %s",
                    (record["record_id"],),
                )
                row = cur.fetchone()
                cols = [desc[0] for desc in cur.description]
                updated = dict(zip(cols, row))
                # File the document (same as ✅ confirm logic)
                review_queue_folder_id = os.environ.get("NOTEBOOKLM_REVIEW_QUEUE_FOLDER_ID", "")
                target_path = updated.get("target_folder_path", "") or ""
                if not target_path.startswith("00_Review_Queue/"):
                    GWSDriveMoveRenameTool()._run(_json.dumps({
                        "file_id": updated["drive_file_id"],
                        "new_name": updated["standardized_filename"],
                        "new_parent_id": review_queue_folder_id,
                        "old_parent_id": os.environ.get("NOTEBOOKLM_LIBRARY_ROOT_ID", ""),
                    }))
                queue_sheet_id = os.environ.get("NOTEBOOKLM_QUEUE_SHEET_ID", "")
                GWSSheetsAppendRowTool()._run(_json.dumps({
                    "spreadsheet_id": queue_sheet_id,
                    "range": "Sheet1!A:I",
                    "values": [[
                        updated["standardized_filename"],
                        "",
                        updated["domain"],
                        updated["project_id"],
                        updated["topic_or_client"],
                        updated["doc_type"],
                        updated["notebook_assignment"],
                        _dt.utcnow().strftime("%Y-%m-%d"),
                        "No",
                    ]],
                }))
                cur.execute(
                    "UPDATE notebooklm_pending_docs SET resolved = true WHERE record_id = %s",
                    (updated["record_id"],),
                )
                conn.commit()
                conn.close()
                _send_emoji(
                    chat_id,
                    f"Created: {new_project_id} | {payload}\nFiled: {updated['standardized_filename']}",
                )
            except Exception as _exc:
                logger.error(f"Emoji handler 🆕 error: {_exc}", exc_info=True)
                _send_emoji(chat_id, "Error processing command. Check logs.")
            return

        # ---- ❌ Flag document for review ----
        if _matched_emoji == "❌":
            try:
                import json as _json
                from datetime import datetime as _dt
                import sys as _sys; _sys.path.insert(0, "/app/orchestrator_skills") if "/app/orchestrator_skills" not in _sys.path else None; from doc_routing.gws_cli_tools import GWSDriveMoveRenameTool, GWSSheetsAppendRowTool
                conn = _db_connect()
                record = _get_pending_doc(conn, reply_msg_id)
                if not record:
                    _send_emoji(chat_id, "No pending document found.")
                    conn.close()
                    return
                review_queue_folder_id = os.environ.get("NOTEBOOKLM_REVIEW_QUEUE_FOLDER_ID", "")
                GWSDriveMoveRenameTool()._run(_json.dumps({
                    "file_id": record["drive_file_id"],
                    "new_name": record["standardized_filename"],
                    "new_parent_id": review_queue_folder_id,
                    "old_parent_id": os.environ.get("NOTEBOOKLM_LIBRARY_ROOT_ID", ""),
                }))
                queue_sheet_id = os.environ.get("NOTEBOOKLM_QUEUE_SHEET_ID", "")
                GWSSheetsAppendRowTool()._run(_json.dumps({
                    "spreadsheet_id": queue_sheet_id,
                    "range": "Flagged Docs!A:E",
                    "values": [[
                        record["original_filename"],
                        record["standardized_filename"],
                        _dt.utcnow().strftime("%Y-%m-%d"),
                        "Operator flagged via Telegram",
                        "No",
                    ]],
                }))
                cur = conn.cursor()
                cur.execute(
                    "UPDATE notebooklm_pending_docs SET resolved = true WHERE record_id = %s",
                    (record["record_id"],),
                )
                conn.commit()
                conn.close()
                _send_emoji(
                    chat_id,
                    f"Flagged: {record['original_filename']}. Moved to Review Queue and logged.",
                )
            except Exception as _exc:
                logger.error(f"Emoji handler ❌ error: {_exc}", exc_info=True)
                _send_emoji(chat_id, "Error processing command. Check logs.")
            return

        # ---- ➕ Approve routing matrix proposal ----
        if _matched_emoji == "➕":
            try:
                import json as _json
                import sys as _sys; _sys.path.insert(0, "/app/orchestrator_skills") if "/app/orchestrator_skills" not in _sys.path else None; from doc_routing.gws_cli_tools import GWSSheetsAppendRowTool
                conn = _db_connect()
                cur = conn.cursor()
                cur.execute(
                    "SELECT * FROM routing_matrix_proposals WHERE status = 'pending' ORDER BY proposed_at DESC LIMIT 1"
                )
                row = cur.fetchone()
                if not row:
                    conn.close()
                    _send_emoji(chat_id, "No pending routing matrix proposals.")
                    return
                cols = [desc[0] for desc in cur.description]
                proposal = dict(zip(cols, row))
                queue_sheet_id = os.environ.get("NOTEBOOKLM_QUEUE_SHEET_ID", "")
                GWSSheetsAppendRowTool()._run(_json.dumps({
                    "spreadsheet_id": queue_sheet_id,
                    "range": "Routing Matrix!A:G",
                    "values": [[
                        "",
                        proposal["signal_keywords"],
                        proposal["suggested_domain"],
                        proposal["suggested_folder"],
                        "",
                        proposal["suggested_notebook"],
                        "operator-approved",
                    ]],
                }))
                cur.execute(
                    "UPDATE routing_matrix_proposals SET status = 'approved' WHERE proposal_id = %s",
                    (proposal["proposal_id"],),
                )
                conn.commit()
                conn.close()
                _send_emoji(
                    chat_id,
                    f"Routing Matrix updated. New row added for: {proposal['signal_keywords']}",
                )
            except Exception as _exc:
                logger.error(f"Emoji handler ➕ error: {_exc}", exc_info=True)
                _send_emoji(chat_id, "Error processing command. Check logs.")
            return

    # /scan-drive — manually trigger Drive watch, processes ALL files in root folder
    if text.lower().strip() == "/scan-drive":
        from notifier import send_message as _send_msg
        _send_msg(chat_id, "Scanning Drive inbox -- will process all unclassified files and report back.")
        def _do_scan():
            try:
                from scheduler import _run_drive_watch
                _run_drive_watch(scan_all=True)
                _send_msg(chat_id, "Drive scan complete.")
            except Exception as _e:
                logger.error(f"/scan-drive error: {_e}")
                _send_msg(chat_id, f"Drive scan failed: {_e}")
        import threading as _threading
        _threading.Thread(target=_do_scan, daemon=True).start()
        return

    # /lessons [task_type] — list recent lessons
    if text.lower().startswith("/lessons"):
        from notifier import send_message as _send_msg
        parts = text.strip().split(maxsplit=1)
        task_type_filter = parts[1].strip() if len(parts) > 1 else None
        from memory import list_lessons
        rows = list_lessons(task_type=task_type_filter, limit=10)
        if not rows:
            _send_msg(chat_id, "No lessons found" + (f" for '{task_type_filter}'" if task_type_filter else "") + ".")
        else:
            lines = ["Recent lessons" + (f" ({task_type_filter})" if task_type_filter else "") + ":"]
            for r in rows:
                sign = "+" if r["learning_type"] != "negative" else "-"
                lines.append(f"[{r['id']}] {sign} [{r['task_type']}] {r['content'][:120]}")
            _send_msg(chat_id, "\n".join(lines))
        return

    # /purge-lesson [id] — mark a lesson as purged
    if text.lower().startswith("/purge-lesson"):
        from notifier import send_message as _send_msg
        parts = text.strip().split(maxsplit=1)
        if len(parts) < 2 or not parts[1].strip().isdigit():
            _send_msg(chat_id, "Usage: /purge-lesson [id]  — get the id from /lessons")
            return
        lesson_id = int(parts[1].strip())
        from memory import purge_lesson
        ok = purge_lesson(lesson_id)
        _send_msg(chat_id, f"Lesson {lesson_id} purged." if ok else f"Lesson {lesson_id} not found.")
        return

    # /status [job_id] — look up a specific job
    if text.lower().startswith("/status"):
        from notifier import send_message as _send_msg
        from memory import get_job
        parts = text.strip().split(maxsplit=1)
        if len(parts) < 2:
            # Show last job if no ID given
            if chat_id in _last_completed_job:
                prior = _last_completed_job[chat_id]
                _send_msg(chat_id, f"Last job: {prior['job_id'][:8]} | {prior['task_type']} | delivered {prior['delivered_at'].strftime('%H:%M')}")
            else:
                _send_msg(chat_id, "No recent jobs. Send /projects to see history.")
            return
        job_id_hint = parts[1].strip()
        job = get_job(job_id_hint)
        if not job:
            _send_msg(chat_id, f"Job '{job_id_hint}' not found.")
            return
        status = job.get("status", "unknown")
        task_type = job.get("task_type", "?")
        created = job.get("created_at", "?")
        result_preview = (job.get("result") or "")[:200]
        msg = f"Job {job_id_hint[:8]}\nStatus: {status}\nType: {task_type}\nStarted: {created}\n\n{result_preview}"
        _send_msg(chat_id, msg)
        return

    # /projects — show recent jobs from this chat session
    if text.lower().startswith("/projects"):
        from notifier import send_message as _send_msg
        try:
            from memory import _pg_conn
            conn = _pg_conn()
            cur = conn.cursor()
            cur.execute("""
                SELECT job_id, task_type, status, created_at
                FROM job_queue
                WHERE from_number = %s
                ORDER BY created_at DESC
                LIMIT 10
            """, (chat_id,))
            rows = cur.fetchall()
            cur.close()
            conn.close()
            if not rows:
                _send_msg(chat_id, "No projects found for this session yet.")
            else:
                lines = ["Recent projects:"]
                for r in rows:
                    jid, jtype, jstatus, jcreated = r
                    ts = str(jcreated)[:16] if jcreated else "?"
                    lines.append(f"{ts} | {jid[:8]} | {jtype or '?'} | {jstatus}")
                _send_msg(chat_id, "\n".join(lines))
        except Exception as e:
            _send_msg(chat_id, f"Could not load projects: {e}")
        return

    # /cost [days] — LLM spend rollup from llm_calls (default last 7 days)
    if text.lower().startswith("/cost"):
        from notifier import send_message as _send_msg
        from memory import _pg_conn
        parts = text.strip().split(maxsplit=1)
        try:
            days = int(parts[1]) if len(parts) > 1 else 7
        except ValueError:
            days = 7
        try:
            conn = _pg_conn()
            cur = conn.cursor()
            cur.execute(
                """
                SELECT
                    COUNT(*) AS total_calls,
                    ROUND(SUM(cost_usd)::numeric, 4) AS total_usd,
                    ROUND(SUM(cost_usd) FILTER (WHERE ts > NOW() - INTERVAL '24 hours')::numeric, 4) AS today_usd,
                    COUNT(DISTINCT council_run_id) FILTER (WHERE council_run_id IS NOT NULL) AS council_runs
                FROM llm_calls
                WHERE ts > NOW() - INTERVAL '%s days'
                """,
                (days,),
            )
            total_calls, total_usd, today_usd, council_runs = cur.fetchone()
            cur.execute(
                """
                SELECT agent_name, COUNT(*) AS calls, ROUND(SUM(cost_usd)::numeric, 4) AS usd
                FROM llm_calls
                WHERE ts > NOW() - INTERVAL '%s days' AND agent_name IS NOT NULL
                GROUP BY agent_name
                ORDER BY usd DESC NULLS LAST
                LIMIT 8
                """,
                (days,),
            )
            top = cur.fetchall()
            cur.close()
            conn.close()
            lines = [
                f"LLM spend (last {days} days):",
                f"  ${total_usd or 0:.4f} total over {total_calls or 0} calls",
                f"  ${today_usd or 0:.4f} in last 24h",
                f"  {council_runs or 0} council runs",
                "",
                "Top agents by spend:",
            ]
            for agent, calls, usd in top:
                lines.append(f"  ${usd or 0:.4f}  {agent}  ({calls} calls)")
            _send_msg(chat_id, "\n".join(lines))
        except Exception as e:
            _send_msg(chat_id, f"Could not load cost: {e}")
        return

    # /switch <project-name> — set active project context
    if text.lower().startswith("/switch"):
        from notifier import send_message as _send_msg
        parts = text.strip().split(maxsplit=1)
        if len(parts) < 2:
            current = _active_project.get(chat_id, "default")
            _send_msg(chat_id, f"Active project: {current}\nUsage: /switch <project-name>")
            return
        project_name = parts[1].strip().lower().replace(" ", "-")
        _active_project[chat_id] = project_name
        _send_msg(chat_id, f"Switched to project: {project_name}\nAll tasks will use this context until you switch again.")
        return

    # ── Praise / Critique Detection ──────────────────────────────────────
    if os.environ.get("MEMORY_LEARNING_ENABLED", "false").lower() == "true":
        from notifier import send_message as _send_msg
        if _is_praise(text) and chat_id in _last_completed_job:
            prior = _last_completed_job[chat_id]
            logger.info(f"Praise detected for job {prior['job_id']} (task_type={prior['task_type']})")
            from memory import extract_and_save_learnings
            threading.Thread(
                target=extract_and_save_learnings,
                args=(prior["task_request"], prior["task_type"], prior["result_summary"], "preference"),
                daemon=True
            ).start()
            _send_msg(chat_id, "Got it — noted as a good pattern.")
            return

        elif _is_feedback_on_prior_job(text, chat_id):
            prior = _last_completed_job[chat_id]
            logger.info(f"Critique detected for job {prior['job_id']} (task_type={prior['task_type']})")
            from memory import extract_negative_lesson
            threading.Thread(
                target=extract_negative_lesson,
                args=(text, prior["task_type"], prior["result_summary"], chat_id),
                daemon=True
            ).start()
            _send_msg(chat_id, "Got it — I'll avoid that next time.")
            return

    job_id = str(uuid.uuid4())
    from notifier import send_briefing, send_message

    # Resolve session key: active project name if set, else chat_id
    active_project = _active_project.get(chat_id)
    session_key = f"{chat_id}:{active_project}" if active_project else chat_id

    # 1. Classify first so the briefing is accurate
    # Shortcut classify runs FIRST — catches task phrases before obvious-chat filter eats them
    _shortcut = _shortcut_classify(text)
    if _shortcut and _shortcut != "memory_capture":
        task_type = _shortcut
        classification = {"task_type": _shortcut, "confidence": 0.95, "is_unknown": False, "has_email_followup": False}
    elif _shortcut == "memory_capture":
        # memory_capture is handled in chat via save_memory tool — treat as chat
        task_type = "chat"
        classification = {"task_type": "chat", "confidence": 0.95, "is_unknown": False, "has_email_followup": False}
    elif _classify_obvious_chat(text):
        task_type = "chat"
        classification = {"task_type": "chat", "confidence": 0.95, "is_unknown": False, "has_email_followup": False}
    else:
        from router import classify_task
        classification = classify_task(text)
        task_type = classification.get("task_type", "unknown")

    # 2. Send briefing (no-op for chat; includes crew + one-line plan for tasks)
    send_briefing(chat_id, task_type, text)

    # 3. Execute
    loop = asyncio.get_running_loop()
    if task_type == "chat":
        # Inject Qdrant context recall for conversational continuity
        enriched_text = text
        try:
            from memory import query_memory
            memories = query_memory(text, top_k=3)
            if memories:
                context_lines = ["[Relevant past context:"]
                for m in memories:
                    ts = m.get("date", "?")
                    summary = m.get("summary", "")[:120]
                    context_lines.append(f"  {ts}: {summary}")
                context_lines.append("]")
                enriched_text = "\n".join(context_lines) + "\n\n" + text
        except Exception:
            pass  # non-fatal — proceed with plain text

        # Chat runs in threadpool — non-blocking
        result = await loop.run_in_executor(None, lambda: run_chat(message=enriched_text, session_key=session_key))
        send_message(chat_id, result["result"])
    else:
        # Complex tasks run in threadpool — non-blocking, avoids freezing the polling loop
        # Use persistent session_key (project-scoped) so conversation history accumulates
        loop.run_in_executor(
            None,
            lambda: _run_background_job(
                task=text,
                from_number=chat_id,
                session_key=session_key,
                job_id=job_id,
                classification=classification,
            )
        )

async def telegram_polling_loop():
    """
    The ultimate fallback: poll for updates instead of waiting for webhooks.
    Bypasses port exposure, SSL, and DNS issues.
    """
    token = os.environ.get("ORCHESTRATOR_TELEGRAM_BOT_TOKEN")
    if not token:
        logger.error("TELEGRAM_POLLING: No token found. Polling disabled.")
        return

    logger.info("TELEGRAM_POLLING: Starting loop...")
    url = f"https://api.telegram.org/bot{token}/getUpdates"
    offset = 0

    # Ensure webhook is cleared so polling works
    webhook_cleared = False
    for attempt in range(3):
        try:
            async with httpx.AsyncClient(timeout=10.0) as client:
                resp = await client.get(f"https://api.telegram.org/bot{token}/deleteWebhook")
                if resp.status_code == 200 and resp.json().get("result"):
                    webhook_cleared = True
                    break
        except Exception as e:
            logger.warning(f"TELEGRAM_POLLING: deleteWebhook attempt {attempt+1} failed: {e}")
        await asyncio.sleep(2)

    if not webhook_cleared:
        logger.error("TELEGRAM_POLLING: Could not clear webhook after 3 attempts. Polling may conflict with webhook delivery.")

    while True:
        try:
            async with httpx.AsyncClient(timeout=30.0) as client:
                resp = await client.get(url, params={"offset": offset, "timeout": 20})
                if resp.status_code == 200:
                    data = resp.json()
                    for update in data.get("result", []):
                        offset = update["update_id"] + 1
                        # Process update without blocking the loop
                        asyncio.create_task(process_telegram_update(update))
                elif resp.status_code == 401:
                    logger.error("TELEGRAM_POLLING: Invalid Token. Stopping.")
                    break
                else:
                    await asyncio.sleep(5)
        except Exception as e:
            logger.error(f"TELEGRAM_POLLING: Error: {e}", exc_info=True)
            await asyncio.sleep(10)

@app.post("/telegram")
async def telegram_webhook(request: Request, background_tasks: BackgroundTasks):
    """
    Keep webhook endpoint as a backup, though polling is now primary.
    """
    try:
        body = await request.json()
        background_tasks.add_task(process_telegram_update, body)
    except Exception:
        pass
    return {"ok": True}


@app.get("/", response_model=StatusResponse)
def status():
    """System status and capability overview."""
    from router import TASK_TYPES
    
    return StatusResponse(
        status="running",
        service="agentsHQ Orchestrator",
        version="2.0.0",
        task_types=list(TASK_TYPES.keys()),
        agents=[
            "planner", "researcher", "copywriter",
            "web_builder", "app_builder", "code_agent",
            "consulting_agent", "social_media_agent",
            "qa_agent", "orchestrator", "agent_creator"
        ],
        uptime_seconds=time.time() - _start_time
    )


@app.get("/health")
def health():
    """Health check endpoint. Called by Docker and n8n."""
    return {
        "status": "ok",
        "timestamp": datetime.utcnow().isoformat(),
        "uptime_seconds": time.time() - _start_time
    }


class HealthReportRequest(BaseModel):
    status: str      # "GREEN", "YELLOW", or "RED"
    report: str      # full Markdown report text
    date: str        # e.g. "2026-04-07"


@app.post("/internal/health-report")
async def receive_health_report(request: HealthReportRequest, req: Request):
    """
    Internal webhook called by the agentsHQ Weekly Health Check remote agent.
    Accepts the report, renders it as HTML, and emails it to all three addresses.
    Protected by X-Internal-Token header — token set via HEALTH_REPORT_TOKEN env var.
    """
    expected_token = os.environ.get("HEALTH_REPORT_TOKEN", "")
    provided_token = req.headers.get("X-Internal-Token", "")
    if not expected_token or provided_token != expected_token:
        raise HTTPException(status_code=401, detail="Unauthorised")

    try:
        from notifier import send_health_check_report
        email_sent = send_health_check_report(request.status, request.report, request.date)
        logger.info(f"Health check report received: status={request.status}, date={request.date}, email_sent={email_sent}")
        return {"ok": True, "email_sent": email_sent}
    except Exception as e:
        logger.error(f"Health report delivery failed: {e}")
        return {"ok": True, "email_sent": False}


@app.post("/run", response_model=AsyncTaskResponse, status_code=202, dependencies=[Depends(verify_api_key)])
async def run_task(request: TaskRequest, background_tasks: BackgroundTasks):
    """
    Main endpoint — async, fire-and-forget.

    Returns 202 immediately after sending a Telegram ack.
    The crew runs in the background; result is pushed to Telegram when done.

    For synchronous callers (Cursor, Claude Code), use POST /run-sync instead.
    """
    logger.info(f"Request from {request.from_number}: {request.task[:100]}...")

    job_id = str(uuid.uuid4())

    # Handle 'more' command synchronously — it's instant
    if request.task.strip().lower() in ("more", "more please", "continue", "show more"):
        from memory import get_next_chunk
        from notifier import send_message
        chunk_result = get_next_chunk(request.session_key)
        if chunk_result["found"]:
            suffix = "\n\n[reply 'more' for the rest]" if chunk_result["has_more"] else "\n\n[end of output]"
            send_message(request.from_number, chunk_result["chunk"] + suffix)
        else:
            send_message(request.from_number, "Nothing more to show — that was the full output.")
        return AsyncTaskResponse(job_id=job_id, status="completed", message="Chunk delivered.")

    # Classify first so the briefing is accurate
    from notifier import send_briefing, send_message
    # Shortcut classify runs FIRST — catches task phrases before obvious-chat filter eats them
    _shortcut = _shortcut_classify(request.task)
    if _shortcut and _shortcut != "memory_capture":
        task_type = _shortcut
        classification = {"task_type": _shortcut, "confidence": 0.95, "is_unknown": False, "has_email_followup": False}
    elif _shortcut == "memory_capture":
        task_type = "chat"
        classification = {"task_type": "chat", "confidence": 0.95, "is_unknown": False, "has_email_followup": False}
    elif _classify_obvious_chat(request.task):
        task_type = "chat"
        classification = {"task_type": "chat", "confidence": 0.95, "is_unknown": False, "has_email_followup": False}
    else:
        from router import classify_task
        classification = classify_task(request.task)
        task_type = classification.get("task_type", "unknown")

    send_briefing(request.from_number, task_type, request.task)

    # For chat, run in executor (non-blocking) and deliver directly
    if task_type == "chat":
        import asyncio
        loop = asyncio.get_running_loop()
        result = await loop.run_in_executor(None, lambda: run_chat(message=request.task, session_key=request.session_key))
        send_message(request.from_number, result["result"])
        return AsyncTaskResponse(job_id=job_id, status="completed", message="Chat response delivered.")

    # Each crew job gets its own session key to prevent concurrent task collision
    crew_session_key = f"{request.session_key}:{job_id[:8]}"

    # Queue crew job in background
    background_tasks.add_task(
        _run_background_job,
        task=request.task,
        from_number=request.from_number,
        session_key=crew_session_key,
        job_id=job_id,
        classification=classification,
    )

    return AsyncTaskResponse(
        job_id=job_id,
        status="pending",
        message="Job queued. Result will be delivered to Telegram.",
    )


@app.post("/run-sync", response_model=TaskResponse, dependencies=[Depends(verify_api_key)])
async def run_task_sync(request: TaskRequest):
    """
    Synchronous endpoint — blocks until result is ready, returns TaskResponse.
    Use this for programmatic callers (Cursor, Claude Code, scripts).
    Does NOT send Telegram messages.
    """
    logger.info(f"[sync] Request from {request.from_number}: {request.task[:100]}...")

    try:
        if request.task.strip().lower() in ("more", "more please", "continue", "show more"):
            from memory import get_next_chunk
            chunk_result = get_next_chunk(request.session_key)
            if chunk_result["found"]:
                suffix = "\n\n[reply 'more' for the rest]" if chunk_result["has_more"] else "\n\n[end of output]"
                reply = chunk_result["chunk"] + suffix
            else:
                reply = "Nothing more to show — that was the full output."
            return TaskResponse(success=True, result=reply, task_type="more", files_created=[], execution_time=0.0)

        shortcut = _shortcut_classify(request.task)
        if shortcut:
            result = run_orchestrator(
                task_request=request.task,
                from_number=request.from_number,
                session_key=request.session_key,
            )
        elif _classify_obvious_chat(request.task):
            result = run_chat(message=request.task, session_key=request.session_key)
        else:
            result = run_orchestrator(
                task_request=request.task,
                from_number=request.from_number,
                session_key=request.session_key,
            )

        return TaskResponse(
            success=result["success"],
            result=result["result"],
            task_type=result.get("task_type", "unknown"),
            files_created=result.get("files_created", []),
            execution_time=result.get("execution_time", 0.0),
            title=result.get("title", ""),
            deliverable=result.get("deliverable", ""),
        )

    except Exception as e:
        logger.error(f"[sync] Request failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Agent execution failed: {str(e)}")


@app.post("/run-team", response_model=TaskResponse, dependencies=[Depends(verify_api_key)])
async def run_team(request: TeamTaskRequest):
    """
    Run multiple crews in parallel (agent teams pattern).

    The caller decomposes the request into independent subtasks.
    All subtasks run concurrently; results are synthesized into one output.

    Example body:
    {
      "original_request": "I need research + a LinkedIn post + a Python script about AI in HR",
      "from_number": "7792432594",
      "subtasks": [
        {"crew_type": "research_crew", "task": "Research AI in HR", "label": "research"},
        {"crew_type": "social_crew",   "task": "Write LinkedIn post on AI in HR", "label": "social"},
        {"crew_type": "code_crew",     "task": "Write CV scorer Python script", "label": "code"}
      ]
    }
    """
    logger.info(f"/run-team received {len(request.subtasks)} subtasks from {request.from_number}")

    if not request.subtasks:
        raise HTTPException(status_code=400, detail="subtasks list cannot be empty")

    try:
        result = run_team_orchestrator(
            subtasks=request.subtasks,
            original_request=request.original_request,
            from_number=request.from_number
        )
        return TaskResponse(
            success=result["success"],
            result=result["result"],
            task_type=result["task_type"],
            files_created=result.get("files_created", []),
            execution_time=result.get("execution_time", 0.0)
        )
    except Exception as e:
        logger.error(f"/run-team failed: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Team execution failed: {str(e)}")


@app.get("/classify", dependencies=[Depends(verify_api_key)])
def classify_only(task: str):
    """
    Classify a task without running it.
    Useful for testing the router and understanding what crew would be used.
    """
    from router import classify_task, get_crew_type
    classification = classify_task(task)
    crew_type = get_crew_type(classification.get("task_type", "unknown"))
    return {
        "task": task,
        "classification": classification,
        "crew_type": crew_type
    }


@app.get("/capabilities", dependencies=[Depends(verify_api_key)])
def capabilities():
    """List all task types the system can handle."""
    from router import TASK_TYPES
    return {
        "task_types": {
            k: {
                "description": v["description"],
                "crew": v["crew"],
                "example_keywords": v["keywords"][:5]
            }
            for k, v in TASK_TYPES.items()
        }
    }


@app.get("/outputs", dependencies=[Depends(verify_api_key)])
def list_outputs():
    """List all files created by the agents."""
    output_dir = os.environ.get("AGENTS_OUTPUT_DIR", "/app/outputs")
    if not os.path.exists(output_dir):
        return {"files": [], "count": 0}
    
    files = []
    for f in os.listdir(output_dir):
        if not f.startswith("."):
            filepath = os.path.join(output_dir, f)
            files.append({
                "name": f,
                "size_bytes": os.path.getsize(filepath),
                "created": datetime.fromtimestamp(
                    os.path.getctime(filepath)
                ).isoformat()
            })
    
    files.sort(key=lambda x: x["created"], reverse=True)
    return {"files": files, "count": len(files)}


@app.get("/outputs/{filename}", dependencies=[Depends(verify_api_key)])
def get_output(filename: str):
    """Retrieve a specific output file."""
    # Security: prevent path traversal
    if ".." in filename or "/" in filename:
        raise HTTPException(status_code=400, detail="Invalid filename")
    
    filepath = f"/app/outputs/{filename}"
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")
    
    with open(filepath, "r", encoding="utf-8") as f:
        content = f.read()
    
    return {
        "filename": filename,
        "content": content,
        "size_bytes": len(content)
    }


@app.get("/memory/search", dependencies=[Depends(verify_api_key)])
def search_memory(query: str, top_k: int = 3):
    """Search agent memory for relevant past tasks."""
    try:
        from memory import query_memory
        results = query_memory(query, top_k=top_k)
        return {"query": query, "results": results, "count": len(results)}
    except Exception as e:
        return {"query": query, "results": [], "error": str(e)}


@app.get("/history/{session_id}", dependencies=[Depends(verify_api_key)])
def get_history(session_id: str, limit: int = 10):
    """Get conversation history for a Telegram session."""
    try:
        from memory import get_conversation_history
        history = get_conversation_history(session_id, limit=limit)
        return {"session_id": session_id, "history": history}
    except Exception as e:
        return {"session_id": session_id, "history": [], "error": str(e)}


@app.post("/run-async", response_model=AsyncTaskResponse, dependencies=[Depends(verify_api_key)])
async def run_task_async(request: TaskRequest, background_tasks: BackgroundTasks):
    """
    Async task endpoint — returns job_id immediately, runs crew in background.

    Use for any task that might take >30s (website builds, app builds, research).
    Poll GET /status/{job_id} until status == 'completed' or 'failed'.

    n8n pattern:
      1. POST /run-async → get job_id (instant)
      2. Send "working on it..." to Telegram
      3. Poll /status/{job_id} every 10s
      4. When completed, send result to Telegram
    """
    job_id = str(uuid.uuid4())[:8]  # short 8-char ID, easy to reference

    from memory import create_job
    create_job(
        job_id=job_id,
        session_key=request.session_key,
        from_number=request.from_number,
        task=request.task
    )

    def _run_in_background():
        from memory import update_job
        try:
            update_job(job_id, status="running")

            # Inject uploaded file content into task if file_id provided
            task_text = request.task
            if request.file_id:
                upload_dir = "/app/uploads"
                matches = [f for f in os.listdir(upload_dir) if f.startswith(request.file_id + "_")] if os.path.isdir(upload_dir) else []
                if matches:
                    txt_path = os.path.join(upload_dir, matches[0] + ".txt")
                    if os.path.exists(txt_path):
                        with open(txt_path, encoding="utf-8", errors="replace") as fh:
                            file_content = fh.read()
                        fname = matches[0][9:]  # strip file_id_ prefix
                        task_text = f"[Attached file: {fname}]\n{file_content}\n\n---\n{request.task}"

            # Use the same classification pipeline as /run and /run-sync
            _shortcut = _shortcut_classify(task_text)
            if _shortcut and _shortcut != "memory_capture":
                _routed_as_chat = False
            elif _classify_obvious_chat(task_text):
                _routed_as_chat = True
            else:
                _routed_as_chat = False

            if _routed_as_chat:
                result = run_chat(message=task_text, session_key=request.session_key)
            else:
                result = run_orchestrator(
                    task_request=task_text,
                    from_number=request.from_number,
                    session_key=request.session_key
                )

            result_text = result.get("result") or result.get("deliverable") or result.get("summary") or result.get("output") or ""
            task_type_val = result.get("task_type", "unknown")
            title_val = result.get("title", task_text[:80])
            deliverable_val = result.get("deliverable", result_text)

            # Save to Drive only for tasks that produce tangible artifacts
            drive_url = ""
            if task_type_val in SAVE_REQUIRED_TASK_TYPES:
                try:
                    from saver import save_to_drive
                    drive_url = save_to_drive(title_val, task_type_val, deliverable_val)
                    if drive_url:
                        result_text = result_text + f"\n\nDrive: {drive_url}"
                except Exception as _drive_err:
                    logger.warning(f"Drive save failed for job {job_id}: {_drive_err}")
            else:
                logger.info(f"Async job {job_id}: task_type '{task_type_val}' is query-only — skipping Drive save")

            update_job(
                job_id=job_id,
                status="completed",
                result=result_text,
                task_type=task_type_val,
                files_created=result.get("files_created", []),
                execution_time=result.get("execution_time", 0.0)
            )
            logger.info(f"Async job {job_id} completed ({result.get('task_type')})")

            # Fire callback if provided
            if request.callback_url:
                try:
                    import requests as _requests
                    _requests.post(request.callback_url, json={
                        "job_id": job_id,
                        "status": "completed",
                        "result": result_text,
                        "task_type": result.get("task_type", "unknown"),
                        "files_created": result.get("files_created", []),
                        "execution_time": result.get("execution_time", 0.0),
                        "from_number": request.from_number,
                        "chat_id": (request.context or {}).get("chat_id", request.from_number),
                        "success": True
                    }, timeout=10)
                    logger.info(f"Callback fired for job {job_id}")
                except Exception as cb_err:
                    logger.warning(f"Callback failed for job {job_id}: {cb_err}")

        except Exception as e:
            logger.error(f"Async job {job_id} failed: {e}", exc_info=True)
            from memory import update_job as uj
            raw = str(e)
            # Sanitize known provider-shaped errors into plain English for the chat UI.
            if "assistant message prefill" in raw or "must end with a user message" in raw:
                friendly = (
                    "The research agent hit a provider limit on this prompt. "
                    "Try asking one focused question at a time (for example, "
                    "'find 5 mechanic shops near 84095 that do safety inspection') "
                    "and I'll run each one cleanly."
                )
            elif "Provider returned error" in raw or "invalid_request_error" in raw:
                friendly = (
                    "The model provider rejected this request. Reword the prompt "
                    "or narrow its scope and try again. I've logged the full trace."
                )
            elif "rate" in raw.lower() and "limit" in raw.lower():
                friendly = "Hit a rate limit. Wait 30 seconds and try again."
            else:
                friendly = f"Task failed. (Diagnostic: {raw[:200]})"
            uj(job_id=job_id, status="failed", error=friendly, result=friendly)

            # Fire failure callback if provided
            if request.callback_url:
                try:
                    import requests as _requests
                    _requests.post(request.callback_url, json={
                        "job_id": job_id,
                        "status": "failed",
                        "result": friendly,
                        "task_type": "unknown",
                        "files_created": [],
                        "execution_time": 0.0,
                        "from_number": request.from_number,
                        "chat_id": (request.context or {}).get("chat_id", request.from_number),
                        "success": False
                    }, timeout=10)
                except Exception:
                    pass

    background_tasks.add_task(_run_in_background)
    logger.info(f"Queued async job {job_id} for: {request.task[:60]}... callback={request.callback_url or 'none'} context={request.context}")

    return AsyncTaskResponse(
        job_id=job_id,
        status="pending",
        message=f"Job {job_id} queued. Poll /status/{job_id} for updates."
    )


@app.get("/status/{job_id}", response_model=JobStatusResponse, dependencies=[Depends(verify_api_key)])
def get_job_status(job_id: str):
    """
    Poll this endpoint after calling /run-async.
    Returns current status: pending | running | completed | failed.
    When completed, result and files_created are populated.
    """
    from memory import get_job
    job = get_job(job_id)

    if not job:
        raise HTTPException(status_code=404, detail=f"Job '{job_id}' not found")

    return JobStatusResponse(
        job_id=job["job_id"],
        status=job["status"],
        task_type=job.get("task_type") or "",
        result=job.get("result") or "",
        files_created=job.get("files_created") or [],
        execution_time=job.get("execution_time") or 0.0,
        error=job.get("error") or ""
    )


class SyncSessionRequest(BaseModel):
    session_key: str                # Telegram chat_id or project-scoped key
    summary: str                    # What happened in the browser session
    source: str = "browser"         # "browser" | "telegram" | "api"
    notify_telegram: bool = False   # If true, send a Telegram notification


@app.post("/inbound-lead", dependencies=[Depends(verify_api_key)])
async def inbound_lead_webhook(request: Request, background_tasks: BackgroundTasks):
    """Webhook for n8n Calendly/Formspree inbound lead events.

    Runs the routine in a background task so n8n gets a fast 202. Result lands
    in Notion Pipeline + Gmail drafts; Telegram notification goes out at the end.
    """
    try:
        body = await request.json()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid JSON body: {exc}")

    logger.info(
        f"Inbound lead received: {body.get('email', '(no email)')} "
        f"source={body.get('source', '?')} booking={body.get('booking_id', '?')}"
    )

    def _run_inbound(payload: dict):
        try:
            from skills.inbound_lead.runner import run_inbound_lead
            from notifier import send_message
            result = run_inbound_lead(payload)
            try:
                from skills.inbound_lead.telegram_notify import format_inbound_telegram_message
                msg = format_inbound_telegram_message(result)
                telegram_chat = os.environ.get("TELEGRAM_CHAT_ID", "")
                if telegram_chat:
                    send_message(telegram_chat, msg)
            except Exception as notify_exc:
                logger.warning(f"Telegram notify failed: {notify_exc}")
            logger.info(
                f"Inbound lead done: status={result.status} "
                f"email={result.payload.email} page={(result.log.notion_page_id if result.log else None)}"
            )
        except Exception as exc:
            logger.error(f"Inbound lead background task failed: {exc}")

    background_tasks.add_task(_run_inbound, body)
    return {"status": "accepted", "message": "Inbound lead queued."}


class ChatTokenRequest(BaseModel):
    pin: str  # simple static PIN gate — keeps crawlers out


@app.post("/chat-token")
async def chat_token(req: ChatTokenRequest):
    """
    Issue a short-lived JWT for the browser chat UI.
    The browser sends this token as Authorization: Bearer <token> on every /run-async call.
    The real ORCHESTRATOR_API_KEY never leaves the server.

    PIN is set via CHAT_UI_PIN env var (required). No PIN = endpoint disabled.
    """
    import jwt as pyjwt
    from datetime import timedelta

    expected_pin = os.environ.get("CHAT_UI_PIN", "")
    if not expected_pin:
        raise HTTPException(status_code=503, detail="Chat UI not configured on this server.")
    if req.pin != expected_pin:
        raise HTTPException(status_code=401, detail="Invalid PIN.")

    secret = os.environ.get("ORCHESTRATOR_API_KEY")
    if not secret:
        logger.error("/chat-token: ORCHESTRATOR_API_KEY not configured -- cannot sign JWT")
        raise HTTPException(status_code=503, detail="Chat token issuer not configured.")
    payload = {
        "sub": "browser-chat",
        "iat": datetime.utcnow(),
        "exp": datetime.utcnow() + timedelta(hours=8),
    }
    token = pyjwt.encode(payload, secret, algorithm="HS256")
    return {"token": token}


def verify_chat_token(authorization: Optional[str] = Header(None)):
    """Dependency: accepts either the raw API key OR a valid browser JWT."""
    import jwt as pyjwt

    if not authorization:
        raise HTTPException(status_code=401, detail="Missing Authorization header.")

    # Allow raw API key (existing integrations)
    api_key = os.environ.get("ORCHESTRATOR_API_KEY", "")
    if authorization == api_key or authorization == f"Bearer {api_key}":
        return True

    # Validate browser JWT
    if authorization.startswith("Bearer "):
        token = authorization[7:]
        try:
            pyjwt.decode(token, api_key, algorithms=["HS256"])
            return True
        except pyjwt.ExpiredSignatureError:
            raise HTTPException(status_code=401, detail="Session expired. Refresh the page to get a new token.")
        except pyjwt.InvalidTokenError:
            raise HTTPException(status_code=401, detail="Invalid token.")

    raise HTTPException(status_code=401, detail="Invalid authorization.")


def _extract_file_text(path: str, filename: str) -> str:
    """Extract readable text from an uploaded file for injection into task prompt."""
    import mimetypes
    ext = os.path.splitext(filename)[1].lower()

    try:
        if ext == ".pdf":
            import pdfplumber
            pages = []
            with pdfplumber.open(path) as pdf:
                for p in pdf.pages[:20]:
                    t = p.extract_text()
                    if t:
                        pages.append(t)
            return "\n\n".join(pages) or "[PDF: no extractable text]"

        if ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip()) or "[DOCX: empty]"

        if ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets[:3]:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in list(sheet.iter_rows(values_only=True))[:100]:
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines) or "[XLSX: empty]"

        if ext == ".csv":
            import csv as _csv
            with open(path, newline="", encoding="utf-8", errors="replace") as f:
                rows = list(_csv.reader(f))[:100]
            return "\n".join(",".join(r) for r in rows) or "[CSV: empty]"

        if ext in (".txt", ".md", ".json", ".py", ".js", ".ts", ".html", ".css", ".yaml", ".yml", ".xml", ".sql"):
            with open(path, encoding="utf-8", errors="replace") as f:
                return f.read(50000) or "[File: empty]"

        if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"):
            from PIL import Image
            img = Image.open(path)
            return f"[Image: {img.format} {img.width}x{img.height}px — vision analysis required]"

        return f"[File: {filename} ({ext}) — binary format, attached for agent reference]"

    except Exception as e:
        logger.warning(f"File extraction failed for {filename}: {e}")
        return f"[File: {filename} — could not extract text: {e}]"


@app.post("/upload", dependencies=[Depends(verify_api_key)])
async def upload_file(file: UploadFile):
    """
    Accept a multipart file upload from the browser chat UI.
    Saves to /app/uploads/, extracts text content, returns file_id.
    The browser passes file_id in the next /run-async call and the orchestrator
    prepends the extracted text to the task prompt.
    """
    import uuid

    upload_dir = "/app/uploads"
    os.makedirs(upload_dir, exist_ok=True)

    file_id = str(uuid.uuid4())[:8]
    safe_name = os.path.basename(file.filename or "upload").replace(" ", "_")
    dest = os.path.join(upload_dir, f"{file_id}_{safe_name}")

    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large (50 MB max).")

    with open(dest, "wb") as f:
        f.write(content)

    extracted = _extract_file_text(dest, safe_name)
    with open(dest + ".txt", "w", encoding="utf-8") as f:
        f.write(extracted)

    logger.info(f"Upload: {safe_name} ({len(content)} bytes) -> {dest}")
    return {
        "file_id": file_id,
        "filename": safe_name,
        "size_bytes": len(content),
        "preview": extracted[:300],
    }


@app.post("/sync-session", dependencies=[Depends(verify_api_key)])
async def sync_session(req: SyncSessionRequest):
    """
    Write a browser session summary into the orchestrator's PostgreSQL conversation
    history under the given session_key. This is Phase 5 Option 1: one-direction bridge.

    thepopebot calls this at the end of a conversation to seed the orchestrator memory
    so Telegram pick-up works with full context.

    No crew is spun up. No LLM call. Just a fast PostgreSQL write + optional Telegram ping.
    """
    from memory import save_conversation_turn
    label = f"[Browser session via {req.source}]"
    content = f"{label}\n{req.summary}"
    save_conversation_turn(req.session_key, "assistant", content)
    save_conversation_turn(req.session_key, "user", f"(Context synced from {req.source} — ready to continue)")

    if req.notify_telegram:
        telegram_chat_id = req.session_key.split(":")[0]
        from notifier import send_message as _send
        _send(telegram_chat_id, f"Browser session saved. Pick up here anytime — I remember what happened.")

    logger.info(f"sync-session: wrote {len(req.summary)} chars for session_key={req.session_key} source={req.source}")
    return {"success": True, "session_key": req.session_key, "chars_written": len(req.summary)}


# ── Run directly for testing ───────────────────────────────────
if __name__ == "__main__":
    import uvicorn
    
    # Ensure required directories exist
    os.makedirs("/app/outputs", exist_ok=True)
    os.makedirs("/app/logs", exist_ok=True)
    os.makedirs("/app/outputs/proposals", exist_ok=True)
    
    uvicorn.run(
        app,
        host="0.0.0.0",
        port=8000,
        log_level="info"
    )
